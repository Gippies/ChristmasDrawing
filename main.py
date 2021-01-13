from openpyxl import load_workbook


def match_people():
    wb = load_workbook("TestFile.xlsx")
    sheet = wb.worksheets[0]
    print(sheet.cell(row=1, column=1).value)


if __name__ == '__main__':
    match_people()
